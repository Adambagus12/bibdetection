import sys
import streamlit as st
import tempfile
import os
import pandas as pd
import subprocess
import cv2
from datetime import datetime
from io import BytesIO
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill

st.set_page_config(
    page_title="Deteksi Bib Pelari",
    page_icon="🏃",
    layout="wide",
    initial_sidebar_state="expanded",
)

# =========================
# CUSTOM STYLE
# =========================
st.markdown(
    """
    <style>
    .block-container {
        padding-top: 2rem;
        padding-bottom: 2rem;
    }
    .section-divider {
        margin-top: 0.5rem;
        margin-bottom: 1.2rem;
        border: none;
        border-top: 1px solid var(--secondary-background-color, rgba(128,128,128,0.3));
        opacity: 0.6;
    }
    div[data-testid="stMetric"] {
        background-color: var(--secondary-background-color, rgba(128,128,128,0.08));
        border: 1px solid var(--secondary-background-color, rgba(128,128,128,0.3));
        border-radius: 10px;
        padding: 12px 16px;
    }
    div[data-testid="stMetric"] label,
    div[data-testid="stMetric"] div {
        color: var(--text-color) !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# =========================
# SETUP ALAMAT ABSOLUT
# =========================
APP_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_OUT_PATH = os.path.join(APP_DIR, "output", "hasil_bib.csv")
FRAME_OUT_PATH = os.path.join(APP_DIR, "output", "frame.jpg")
SCRIPT_MAIN_PATH = os.path.join(APP_DIR, "main.py")

# =========================
# KONFIGURASI SISTEM (TETAP, TIDAK DITAMPILKAN DI UI)
# =========================
# Nilai berikut sudah dikunci berdasarkan pengujian sistematis pada Bab 3
# dan Bab 4 (confidence 0.45 dipilih setelah dibandingkan dengan 0.25 dan
# 0.65; frame_skip = 1 berarti YOLOv8n mendeteksi di SETIAP frame, karena
# penyaringan frekuensi pemrosesan sudah dilakukan di tahap OCR internal
# main.py, bukan di tahap deteksi). Operator tidak dapat mengubah nilai ini
# dari antarmuka.
CONFIDENCE = 0.45
FRAME_SKIP = 1
AUTO_FLAG = "0"

# =========================
# SESSION STATE
# =========================
defaults = {
    "logs": [],
    "detections": [],
    "processing": False,
    "done": False,
    "video_server_path": None,
    "video_uploaded": False,
}
for key, val in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = val


def add_log(msg: str):
    ts = datetime.now().strftime("%H:%M:%S")
    st.session_state.logs.append(f"[{ts}] {msg}")


def add_detection(bib: str, frame_num: str, total: str):
    ts = datetime.now().strftime("%H:%M:%S")
    no = len(st.session_state.detections) + 1
    st.session_state.detections.append(f"#{no:03d} [{ts}] Bib {bib} — frame {frame_num}/{total}")


# =========================
# SIDEBAR — INFO & KONFIGURASI
# =========================
with st.sidebar:
    st.markdown("### 🏃 Deteksi Bib Pelari")
    st.caption("Sistem deteksi & pengenalan nomor bib berbasis YOLOv8n + EasyOCR")

    st.markdown("---")
    st.markdown("**Status Berkas**")
    video_ready = st.session_state.video_uploaded
    st.write("🎬 Video:", "✅ Siap" if video_ready else "⏳ Belum diunggah")

# =========================
# HEADER UTAMA
# =========================
st.title("Sistem Deteksi dan Pengenalan Nomor Bib Pelari")
st.caption("Unggah video lomba lari dan database peserta, lalu jalankan proses deteksi secara otomatis.")
st.markdown('<hr class="section-divider">', unsafe_allow_html=True)

# =========================
# TAHAP 1 — UPLOAD
# =========================
st.subheader("1. Unggah Berkas")

col1, col2 = st.columns([1, 1], gap="large")

with col1:
    video_file = st.file_uploader("🎬 Video Lomba Lari", type=["mp4"])
    db_file = st.file_uploader("📂 Database Peserta (.csv / .xlsx)", type=["csv", "xlsx"])

    if video_file:
        size_mb = round(video_file.size / 1024 / 1024, 2)
        st.info(f"📁 {video_file.name} · {size_mb} MB")

        if st.session_state.get("last_video_name") != video_file.name:
            st.session_state.video_uploaded = False
            st.session_state.video_server_path = None
            st.session_state["last_video_name"] = video_file.name

        if not st.session_state.video_uploaded:
            save_bar = st.progress(0, text="💾 Menyimpan video...")
            try:
                video_data = video_file.read()
                total_size = len(video_data)
                chunk_size = 1024 * 512
                uploaded = 0

                temp_dir = tempfile.mkdtemp()
                local_path = os.path.join(temp_dir, video_file.name)

                with open(local_path, "wb") as f:
                    for i in range(0, total_size, chunk_size):
                        chunk = video_data[i:i + chunk_size]
                        f.write(chunk)
                        uploaded += len(chunk)
                        pct = min(int(uploaded / total_size * 100), 100)
                        save_bar.progress(pct, text=f"💾 Menyimpan video... {pct}%")

                st.session_state.video_server_path = local_path
                st.session_state.video_uploaded = True
                save_bar.progress(100, text="✅ Video siap diproses!")
                add_log(f"Video disimpan: {video_file.name} ({size_mb} MB)")
            except Exception as e:
                save_bar.empty()
                st.error(f"❌ Gagal menyimpan: {str(e)}")
        else:
            st.success("✅ Video siap diproses")

    if db_file:
        db_valid = True
        db_error_msg = ""
        try:
            ext = os.path.splitext(db_file.name)[1].lower()
            if ext == ".csv":
                df_check = pd.read_csv(db_file)
            elif ext in (".xlsx", ".xls"):
                df_check = pd.read_excel(db_file)
            else:
                db_valid = False
                db_error_msg = f"Format '{ext}' tidak didukung. Gunakan .csv atau .xlsx."

            if db_valid:
                df_check.columns = [c.strip() for c in df_check.columns]
                kolom_wajib = ["BIB", "Nama", "Kategori"]
                kolom_hilang = [k for k in kolom_wajib if k not in df_check.columns]
                if kolom_hilang:
                    db_valid = False
                    db_error_msg = (
                        f"Kolom wajib tidak ditemukan: {', '.join(kolom_hilang)}. "
                        f"Pastikan berkas memiliki kolom 'BIB', 'Nama', dan 'Kategori'."
                    )
        except Exception as e:
            db_valid = False
            db_error_msg = f"Gagal membaca berkas: {str(e)}"
        finally:
            db_file.seek(0)  # reset pointer supaya bisa dibaca ulang saat proses dijalankan

        st.session_state["db_valid"] = db_valid

        if db_valid:
            st.success(f"✅ {db_file.name}")
            add_log(f"Database dipilih: {db_file.name}")
        else:
            st.error(f"❌ Database tidak valid: {db_error_msg}")
            add_log(f"❌ Database tidak valid: {db_error_msg}")

with col2:
    if video_file:
        st.markdown("**Pratinjau Video**")
        st.video(video_file)
    else:
        st.info("Pratinjau video akan muncul di sini setelah diunggah.")

st.markdown('<hr class="section-divider">', unsafe_allow_html=True)

# =========================
# TAHAP 2 — PROSES
# =========================
st.subheader("2. Jalankan Proses")

process_clicked = st.button(
    "🚀 Mulai Proses Deteksi",
    disabled=st.session_state.processing,
    type="primary",
    use_container_width=True,
)

if process_clicked:
    if not video_file or not db_file:
        st.error("❌ Harap unggah Video dan Database terlebih dahulu.")
        st.stop()


    if not st.session_state.get("db_valid", True):
        st.error("❌ Database tidak valid, silakan unggah ulang berkas database yang benar.")
        st.stop()

    st.session_state.done = False
    st.session_state.detections = []
    st.session_state.processing = True

    add_log("Menyimpan file sementara...")

    temp_dir = tempfile.mkdtemp()
    video_path = st.session_state.video_server_path

    db_path = os.path.join(temp_dir, db_file.name)
    with open(db_path, "wb") as f:
        f.write(db_file.read())
    add_log(f"Database disimpan: {db_file.name}")
    add_log(f"Mulai proses — CONF:{CONFIDENCE} SKIP:{FRAME_SKIP} AUTO:{AUTO_FLAG}")

    progress_bar = st.progress(0, text="Memproses video...")

    col_preview, col_det = st.columns([1, 1], gap="large")
    with col_preview:
        st.markdown("**📸 Pratinjau Langsung**")
        frame_placeholder = st.empty()
    with col_det:
        det_header = st.empty()
        det_box = st.empty()

    with st.expander("📋 Log Proses", expanded=False):
        log_box = st.empty()

    def render_panels():
        total_det = len(st.session_state.detections)
        det_header.markdown(f"**🎯 Bib Terdeteksi ({total_det})**")
        det_box.code(
            "\n".join(st.session_state.detections[-40:]) if st.session_state.detections else "— belum ada deteksi —",
            language="bash"
        )
        log_box.code(
            "\n".join(st.session_state.logs[-50:]) if st.session_state.logs else "— belum ada aktivitas —",
            language="bash"
        )

    render_panels()

    try:
        process = subprocess.Popen(
            [sys.executable, SCRIPT_MAIN_PATH, video_path, db_path, str(CONFIDENCE), str(FRAME_SKIP), AUTO_FLAG],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            encoding="utf-8",
            errors="ignore"
        )

        while True:
            line = process.stdout.readline()
            if not line:
                break
            line = line.strip()
            if not line:
                continue

            if line.startswith("PROGRESS:"):
                try:
                    p = max(0, min(int(line.split(":")[1].strip()), 100))
                    progress_bar.progress(p, text=f"Memproses video... {p}%")
                except Exception:
                    pass

            elif line.startswith("DETECTED:"):
                try:
                    parts = line.split(":")
                    bib = parts[1].strip()
                    frame_num = parts[3].strip() if len(parts) >= 4 else "?"
                    total = parts[5].strip() if len(parts) >= 6 else "?"
                    add_detection(bib, frame_num, total)
                except Exception:
                    pass
                render_panels()

            else:
                add_log(line)
                render_panels()

            if os.path.exists(FRAME_OUT_PATH):
                img = cv2.imread(FRAME_OUT_PATH)
                if img is not None:
                    img = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
                    frame_placeholder.image(img, use_container_width=True)

        process.wait()

        if process.returncode == 0:
            progress_bar.progress(100, text="✅ Selesai!")
            add_log("✅ Proses selesai!")
            st.session_state.done = True
        else:
            progress_bar.progress(100, text="❌ Terjadi error")
            add_log(f"❌ Proses gagal. Return code: {process.returncode}")

    except Exception as e:
        add_log(f"❌ Exception: {str(e)}")
    finally:
        st.session_state.processing = False
        render_panels()

    st.rerun()

st.markdown('<hr class="section-divider">', unsafe_allow_html=True)

# =========================
# TAHAP 3 — HASIL
# =========================
st.subheader("3. Hasil Deteksi")


@st.fragment
def render_hasil():
    if st.session_state.done:
        if os.path.exists(CSV_OUT_PATH) and os.path.getsize(CSV_OUT_PATH) > 0:
            try:
                df = pd.read_csv(CSV_OUT_PATH)
            except pd.errors.EmptyDataError:
                df = pd.DataFrame()

            if df.empty:
                st.warning("⚠️ Tidak ada bib yang berhasil terverifikasi pada video ini.")
            else:
                m1, m2, m3 = st.columns(3)
                m1.metric("Runner Terverifikasi", len(df))
                m2.metric("Total Deteksi Bib", len(st.session_state.detections))
                m3.metric("Status", "Selesai ✅")

                st.markdown("**Tabel Hasil**")
                st.dataframe(df, use_container_width=True)

                video_label = st.session_state.get("last_video_name", "video")
                video_label = os.path.splitext(video_label)[0]
                video_label = "".join(
                    c if c.isalnum() or c in ("-", "_") else "_" for c in video_label
                )
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

                # --- Download Excel (.xlsx) rapi & terformat ---
                df_export = df.copy()
                df_export.insert(0, "No", range(1, len(df_export) + 1))

                excel_buffer = BytesIO()
                with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
                    df_export.to_excel(writer, index=False, sheet_name="Hasil Deteksi")
                    worksheet = writer.sheets["Hasil Deteksi"]

                    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                    header_font = Font(color="FFFFFF", bold=True)
                    for col_idx, column in enumerate(df_export.columns, start=1):
                        cell = worksheet.cell(row=1, column=col_idx)
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = Alignment(horizontal="center", vertical="center")

                    for col_idx, column in enumerate(df_export.columns, start=1):
                        max_len = max(
                            [len(str(column))] + [len(str(v)) for v in df_export[column].astype(str)]
                        )
                        worksheet.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

                    worksheet.freeze_panes = "A2"

                excel_buffer.seek(0)
                xlsx_filename = f"hasil_bib_{video_label}_{timestamp}.xlsx"

                csv_filename = f"hasil_bib_{video_label}_{timestamp}.csv"
                csv_data = df_export.to_csv(index=False).encode("utf-8-sig")

                col_dl1, col_dl2 = st.columns(2)
                with col_dl1:
                    st.download_button(
                        "📄 Download CSV",
                        data=csv_data,
                        file_name=csv_filename,
                        mime="text/csv",
                        use_container_width=True,
                    )
                with col_dl2:
                    st.download_button(
                        "📊 Download Excel",
                        data=excel_buffer,
                        file_name=xlsx_filename,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                    )

                with st.expander(f"🎯 Rincian Bib Terdeteksi ({len(st.session_state.detections)})"):
                    st.code(
                        "\n".join(st.session_state.detections) if st.session_state.detections else "— tidak ada deteksi —",
                        language="bash"
                    )
        else:
            st.warning("⚠️ File hasil tidak ditemukan atau proses belum menghasilkan bib terverifikasi. Path: " + CSV_OUT_PATH)
    else:
        st.info("Hasil akan ditampilkan di sini setelah proses deteksi selesai dijalankan.")


render_hasil()

# =========================
# LOG SISTEM (RIWAYAT PENUH)
# =========================
with st.expander("📋 Log Sistem Lengkap", expanded=False):
    if st.session_state.logs:
        st.code("\n".join(st.session_state.logs), language="bash")
    else:
        st.caption("— belum ada aktivitas —")
    if st.button("🗑️ Hapus Log"):
        st.session_state.logs = []
        st.rerun()